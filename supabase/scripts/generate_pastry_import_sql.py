from __future__ import annotations

import json
import re
import unicodedata
from collections import defaultdict
from pathlib import Path

from docx import Document


DOWNLOADS = Path.home() / "Downloads"
OUTPUT_DIR = Path(__file__).resolve().parent
CODIFICATION_FILE = DOWNLOADS / "codification des produits-3.docx"

RECIPE_FILES = {
    "BROWNIE NOISETTE.docx": "cake",
    "COOKIES AMR PISTACHE.docx": "cake",
    "CAKE GR CHOCO VANILLE PROD.docx": "cake",
    "les glacages.docx": "patisserie",
    "Paris breste.docx": "patisserie",
    "fiche prod les tartes.docx": "cake",
    "POP CAKE.docx": "cake",
    "tromp leile.docx": "patisserie",
    "Fiche de production des cremes.docx": "patisserie",
    "MILLE FEILLE VANILLE CHOCO.docx": "patisserie",
    "RED VEL VET.docx": "patisserie",
    "cake americain.docx": "cake",
    "Noisella.docx": "patisserie",
    "MACARON PRODUCTION.docx": "patisserie",
    "Soiree.docx": "patisserie",
    "madelin-financier.docx": "cake",
    "fich prod fondo chocolat.docx": "cake",
    "FICHE PROD BROWNIE NOIX.docx": "cake",
    "FICHE PRODUCTION PATE SABLE.docx": "patisserie",
    "PRODUCTION FUILLTAGE.docx": "viennoiserie",
    "tarte amande noix pomme.docx": "cake",
    "Bombe d amoure.docx": "patisserie",
    "CAnnelle.docx": "patisserie",
    "production cake dattes.docx": "cake",
    "les eclaire.docx": "patisserie",
}


def clean(value: str) -> str:
    return " ".join(value.replace("\xa0", " ").split()).strip()


def key(value: str) -> str:
    normalized = unicodedata.normalize("NFD", clean(value))
    normalized = "".join(char for char in normalized if unicodedata.category(char) != "Mn")
    normalized = normalized.replace("œ", "oe").replace("Œ", "OE")
    return re.sub(r"[^A-Z0-9]+", " ", normalized.upper()).strip()


ALIASES = {
    "APPAREILLE CAKE GRAND CHOCOLAT": "Appareille cake gr chocolat",
    "APPAREILLE CAKE GR CHOCO": "Appareille cake gr chocolat",
    "APPAREILLE CAKE GRAND VANILLE": "Appareille cake gr vanille",
    "APPARAILLE CAKE GR CHOCOLAT": "Appareille cake gr chocolat",
    "APPARAILLE CAKE GR VANILLE": "Appareille cake gr vanille",
    "APPARAILLE CAKE ROYALE": "Appareille cake royale",
    "APPAREIL CAKE GR CHOCOLAT": "Appareille cake gr chocolat",
    "APPAREIL CAKE GR VANILLE": "Appareille cake gr vanille",
    "APPAREIL CAKE ROYALE": "Appareille cake royale",
    "APPAREILLE COOKIES GR CHOCOLAT": "Pate cookies gr chocolat",
    "APPAREILLE COOKIES GR FRAMBOISE": "Pate cookies gr framboise",
    "APPAREILLE COOKIES GR NOISETTE": "Pate cookies gr noisette",
    "APPAREILLE COOKIES GR PISTACHE": "Pate cookies gr pistache",
    "PATE COOKIES GRAND CHOCOLAT": "Pate cookies gr chocolat",
    "PATE COOKIES GRAND FRAMBOISE": "Pate cookies gr framboise",
    "PATE COOKIES GRAND NOISETTE": "Pate cookies gr noisette",
    "PATE COOKIES GRAND PISTACHE": "Pate cookies gr pistache",
    "COOKIERS GR CHOCOLAT": "Cookies gr chocolat",
    "COOKIERS GR FRAMBOISE": "Cookies gr framboise",
    "COOKIERS GR NOISETTE": "Cookies gr noisette",
    "COOKIERS GR PISTACHE": "Cookies gr pistache",
    "COOKIES FRAMOISE GR": "Cookies gr framboise",
    "COOKIES NOISSETE GR": "Cookies gr noisette",
    "COOKIES NOISETTE GR": "Cookies gr noisette",
    "COOKIES PISTACH GR": "Cookies gr pistache",
    "BROWNIES NOISETTE": "Brownie noisette",
    "BROWNIES NOIX": "Brownie noix",
    "BISCUIT CAKE AMERICAN": "Biscuit Cake Américain",
    "BISCUIT CAKE CAROTTE": "Biscuit cake Amr carotte",
    "BISCUIT CAROTTE": "Biscuit cake Amr carotte",
    "BISCUIT TROIS CHOCOLAT": "Biscuit 3chocolat",
    "CARAMEL BEURRE SALE": "Caramel au beurre salé",
    "CADRE TROIS CHOCOLATS": "Cadre Trois chocolat",
    "CADRE TROPICALE": "Cadre Tropical",
    "CHOCOLAT ROCHEE NOIR": "Chocolat rochée noire",
    "CHOCOLAT ROCHE NOIRE": "Chocolat rochée noire",
    "CHOCOLAT ROCHE": "Chocolat rochée au lait",
    "CHOCOLAT ROCHER NOIRE": "Chocolat rochée noire",
    "CHOCOLAT ROCHER ROUGE": "Chocolat rochée rouge",
    "CHOCOLAT ROCHER": "Chocolat rochée au lait",
    "CREME CHESS": "Crème Cheese",
    "CREME AU BEURRE": "Crème au beurre",
    "CREME BEURRE": "Crème au beurre",
    "CREME PATISSIERE": "Crème pâtissière",
    "CREME TARTE CITRON": "Crème tarte citron",
    "CREMES TARTE CITRON": "Crème tarte citron",
    "CREME MOUSSELINE PISTCHE": "Crème mousseline pistache",
    "CREME MOUSSELINE PRALINE": "Crème mousseline praliné",
    "CREMEUX NOIRE": "Crémeux noir",
    "CREMEUX NOISETTE": "Crémeux noisiola",
    "CREMEUX MANGE": "Crémeux mangue",
    "GLACAGES AU LAIT": "Glaçage au lait",
    "GLACAGES CARAMEL": "Glaçage caramel",
    "GLACAGES MANGE": "Glaçage mangue",
    "GLACAGE MANGE": "Glaçage mangue",
    "GLACAGES OPERA": "Glaçage opéra",
    "GLACAGES PASSION": "Glaçage passion",
    "GLACAGES VANILLE": "Glaçage vanille",
    "GLACAGE NOIRE": "Glaçage noire",
    "COULIS YOGHOURT": "Coulis yaourt",
    "COULIS TARTE SOLEIL": "Coulis soleil",
    "COULIS TROMPE L OEIL MANGE PASSION": "Coulis trompe l’œil mangue passion",
    "COLIS MANGE PASSION": "Coulis mangue passion",
    "COLIS TROPICAL": "Coulis tropical",
    "COLIS ENTREMET FRAMBOISE": "Coulis entremet framboise",
    "COLIS MANGUE PASSION": "Coulis mangue passion",
    "COMPORTE FRAMBOISE": "Compotée framboise",
    "LA MOUSSE CAFE": "Mousse café",
    "LA MOUSSE CASABIANCA": "Mousse casabianca",
    "LA MOUSSE CHEESECAKE MANGUE": "Mousse cheesecake mangue",
    "LA MOUSSE CHEESECAKE MYRTILLE": "Mousse cheesecake myrtille",
    "LA MOUSSE ENTREMET FRAMBOISE": "Mousse entremet framboise",
    "LA MOUSSE ENTREMET MANGE PASSION": "Mousse entremet mangue passion",
    "LA MOUSSE ENTREMET NOISETTE": "Mousse noisette",
    "LA MOUSSE FRAMBOISE": "Mousse framboise",
    "LA MOUSSE FROMAGE CHEESECAKE MYRTILLE": "Mousse fromage cheesecake myrtille",
    "LA MOUSSE MANGUE": "Mousse mangue",
    "LA MOUSSE PINACOLADA": "Mousse pinacolada",
    "LA MOUSSE PISTACHE": "Mousse pistache",
    "LA MOUSSE PRALINEE": "Mousse pralinée",
    "MOUSSE NOIR": "Mousse noire",
    "MOUSSE MANGUE PASSION": "Mousse entremet mangue passion",
    "MOUSSE MANGE": "Mousse mangue",
    "MOUSSE MANGE PASSION": "Mousse entremet mangue passion",
    "MOUSSE COCE": "Mousse coco",
    "L INSERT FRAMBOISES": "L’insert framboise",
    "L INCERT CITRON": "L’insert citron",
    "PATE A CHOUX": "Pate choux",
    "PATE CHOW": "Pate choux",
    "PATE FUILLTAGE": "Pate feuilletage",
    "PATE FEILLTAGE": "Pate feuilletage",
    "PATE GLACE": "Pate glaces noire",
    "PATE GLACE BLANC": "Pate glaces au lait",
    "PATE FEUILLAGE": "Pate feuilletage",
    "PATE ECLAIRE": "Pate choux",
    "PISTOLET CHOCOLAT NOIR": "Pistolet chocolat noire",
    "PISTOLET CHOCOLAT CARAMEL": "Pistolet caramel",
    "PISTOLET VER": "Pistolet vert",
    "PISTOLET VERS": "Pistolet vert",
    "PRALINEE AMANDE": "Praline amande",
    "PRALINEE AMANDE NOIRE": "Praline amande",
    "PRALINEE ARACHIDE": "Praline arachide",
    "PRALINEE NOISETTE": "Praline noisette",
    "PRALINEE PISTACHE": "Praline pistache",
    "PRALINE AMANDE NOIRE": "Praline amande",
    "PRALINER PISTACHE": "Praline pistache",
    "PRALINER": "Praline",
    "PRALINEE": "Praline",
    "PATE D AMANDE": "Pate amande",
    "SIROP SUCRE": "Sirop de sucre",
    "TROPICALE": "Tropical",
    "TROIS CHOCOLATS": "3 Chocolat",
    "TABLETTE KONAFA PISTACHE": "Tablette kounafa",
    "CHAUSSON JALOSIE AMANDE": "Chausson jalousie amande",
    "PALMIE BAGUETTE SUCRE": "Palmier baguette sucre",
    "MACARON PETIT BLEU": "Macaron petit bleu",
    "MACARON PETIT CHOCOLAT NOIRE": "Macaron petit noir",
    "MACARON VIOLET": "Macaron petit violet",
    "ENTREMET CAKE AMERICAN NOUGAT": "Entremet cake American NOUGA",
    "ENTR CAKE AMR FERERO": "Entremet cake American Ferrero",
    "ENTREMET RED VELVET": "Cadre Red velvet",
    "GLACAGE CHEESE CAKE MANGUE": "Glaçage cheesecake Mangue",
    "SILICONE CASABIANCA": "Silicone gâteau Casabianca",
    "LA MOUSSE CHEESE CAKE MANGUE": "Mousse cheesecake mangue",
    "LA MOUSSE CHEESE CAKE MYRTILLE": "Mousse cheesecake myrtille",
    "LA MOUSSE FROMAGE CHEESE CAKE MYRTILLE": "Mousse fromage cheesecake myrtille",
    "NOUGAT": "NOUGA",
    "NOUGAT AMANDE": "Nouga amande",
    "PISTACHE PRALINE": "Praline pistache",
    "PISTOLET C CARAMEL": "Pistolet caramel",
    "ROCHEE NOIRE": "Chocolat rochée noire",
}

FORCED_SEMI_FINISHED = {
    "Biscuit",
    "Croquant",
    "Crème citron",
    "Crème framboise",
    "Ganache",
    "Mousse",
    "Pate bombe",
    "Pate dattes",
    "Pate glaces au lait",
    "Praline",
    "NOUGA",
}

RAW_ALIASES = {
    "AGAR AGAR": "Agar-agar",
    "AGR GAR": "Agar-agar",
    "AGRAR": "Agar-agar",
    "AGR GAR": "Agar-agar",
    "AGRAGAR": "Agar-agar",
    "AGRGR": "Agar-agar",
    "AGRGAR": "Agar-agar",
    "AMANDE EFFILE": "Amande effilée",
    "AMANDE EFFILEE": "Amande effilée",
    "AMANDE HACHEE": "Amande hachée",
    "AMANDE POUDRE": "Amande poudre",
    "BEURRE SPECIALE": "Beurre spécial",
    "BEURRE SPECIAL": "Beurre spécial",
    "B SPECIAL": "Beurre spécial",
    "BUERRE SPECIALE": "Beurre spécial",
    "BEURRE CACAO": "Beurre de cacao",
    "BEURRE BONNA": "Beurre Bonna",
    "BEURRE HUILE": "Huile",
    "HUILE JETABLE BEURRE SPECIALE": "Huile",
    "BLANC D OEUF": "Blanc d’œuf",
    "C FRAICHE": "Crème fraîche",
    "CREME FRICHE": "Crème fraîche",
    "CREME FRAICHE": "Crème fraîche",
    "COULEURENT ROUGE": "Colorant rouge",
    "COULERENT ROUGE": "Colorant rouge",
    "COULERENT JAUNE": "Colorant jaune",
    "COULERENT VERS": "Colorant vert",
    "COULEUR ROUGE": "Colorant rouge",
    "COLORENT VERT": "Colorant vert",
    "COULEURENT": "Colorant",
    "COULEUR": "Colorant",
    "COULERENT NOIRE": "Colorant noir",
    "EAU": "Eau",
    "FARINE FORCE": "Farine force",
    "FEUILLE GELATINE": "Gélatine feuille",
    "GELATINE FEUILLE": "Gélatine feuille",
    "GELATINE MASSE": "Masse gélatine",
    "GELATINE": "Gélatine",
    "HUILE DE TABLE": "Huile",
    "HUILE JETABLE": "Huile",
    "JAUNE D OEUF": "Jaune d’œuf",
    "JAUN OEUF": "Jaune d’œuf",
    "JAUNE OEUF": "Jaune d’œuf",
    "LA FARINE": "Farine",
    "FARINE VIENNOISERIE": "Farine viennoiserie",
    "LEVURE CHIMIQUE": "Levure chimique",
    "MAIZINA": "Maïzena",
    "OEUF": "Œufs",
    "OEUFS": "Œufs",
    "OUEF": "Œufs",
    "PUREE MANGE": "Purée mangue",
    "MANGE": "Mangue",
    "MANGUE PUREE": "Purée mangue",
    "PUREE ANANAS": "Purée ananas",
    "PUREE DE FRAMBOISE": "Purée framboise",
    "PUREE FRAMBOISE": "Purée framboise",
    "PUREE DE CITRON": "Purée citron",
    "PUREE PASSION": "Purée passion",
    "PUREE COCE": "Purée coco",
    "SUCRE GLACE": "Sucre glacé",
    "SUCRE GLACER": "Sucre glacé",
    "SUCRE SEMOLE": "Sucre semoule",
    "VANILLE POUDRE": "Vanille poudre",
    "L AROME CITRON": "Arôme citron",
    "L AROME PISTACHE": "Arôme pistache",
    "TRABLET CAFE": "Trablit café",
    "TRIMOLINE": "Trimoline",
    "ROYALTINE": "Royal tine",
    "XANTHAN": "Xanthane",
    "XANTHINE": "Xanthane",
    "ANANAS CONSERVE": "Ananas conserve",
    "CONSERVE ANANAS": "Ananas conserve",
    "BANANE": "Banane",
    "CHOCO NOIRE": "Chocolat noir",
    "CHOCOLAT NOIRE": "Chocolat noir",
    "CHOCOLAT LAIT": "Chocolat au lait",
    "CHOCOLAT NOIRE SPEC": "Chocolat noir spécial",
    "CHOCOLAT NOIRE SPECIAL": "Chocolat noir spécial",
    "CHOCOLAT DROPS": "Drops",
    "CITRON JUS": "Jus de citron",
    "JUS CITRON": "Jus de citron",
    "FEUILLETINE": "Feuilletine",
    "MANGUE": "Mangue",
    "MIEL": "Miel",
    "NOISETTE": "Noisette",
    "NOISETTE AMANDE NOIRE": "Noisette",
    "NAPPAGE NUDER": "Nappage neutre",
    "POUDRE AMANDE": "Amande poudre",
    "POUDRE CREME": "Poudre pâtissière",
    "POUDRE PATISSERIE": "Poudre pâtissière",
    "POUDRE PATISSIERE": "Poudre pâtissière",
    "GOUSSE VANILLE": "Gousse de vanille",
    "VANILLE GOUSSE": "Gousse de vanille",
}

TYPE_OVERRIDES = {
    "SILICONE PINACOLADA": "semi_finished",
    "SOIREE BLANC": "finished",
    "SOIREE CITRON": "finished",
    "SOIREE NOIRE": "finished",
}

CATEGORY_OVERRIDES = {
    "PATE MF": "viennoiserie",
    "PATE FEUILLETAGE": "viennoiserie",
    "CREME AMANDE": "viennoiserie",
}

OFFICIAL_CODE_OVERRIDES = {
    "APPAREILLE CAKE GR CHOCOLAT": "ACGC",
    "APPAREILLE CAKE GR VANILLE": "ACGV",
    "APPAREILLE CAKE ROYALE": "ACR",
    "BISCUIT CAKE AMERICAIN": "BCA",
    "CARAMEL AU BEURRE SALE": "CBS",
    "CROQUANT PRALINEE": "CQPR",
    "ENTREMET CAKE AMERICAN FERRERO": "23-1-SF",
    "ENTREMET CAKE AMERICAN NOUGA": "23-3-SF",
    "ENTREMET CAKE AMERICAN NUTELLA": "23-2-SF",
    "ENTREMET CAKE AMERICAN CAROTTE": "25SF",
    "GLACAGE CHEESECAKE MANGUE": "GCCM",
    "GLACAGE MANGUE": "GMA",
    "GLACAGE VANILLE": "GV",
    "BROWNIE NOISETTE": "32",
    "BROWNIE NOIX": "31",
    "CAKE AMR CAROTTE": "25",
    "CAKE AMR FERRERO": "23",
    "CAKE AMR NOUGAT": "23",
    "CAKE AMR NUTELLA": "23",
    "ENTREMET FERRERO": "23",
    "COOKIES AMERICAIN": "27-1",
    "COOKIES GR CHOCOLAT": "27-2",
    "COOKIES GR FRAMBOISE": "27-2",
    "COOKIES GR NOISETTE": "27-2",
    "COOKIES GR PISTACHE": "27-2",
    "FINANCI AMANDE": "28",
    "FINANCI PISTACHE": "28",
    "FINANCIE TIGRE": "28",
    "MADELEINE VANILLE": "29",
    "MADELEINE CHOCOLAT": "29",
    "POP CAKE": "36",
    "TARTE AMANDE": "37",
    "TARTE NOIX": "37",
    "TARTE POMME": "37",
    "TARTE CITRON": "37",
    "TARTE FRUIT": "37",
    "TARTE PARIS BREST": "37",
    "TARTE SOLEIL": "37",
    "MACARON PETIT ORANGE": "39",
    "MACARON PETIT BLANC": "39",
    "MACARON PETIT NOIR": "39",
    "MACARON PETIT ROUGE": "39",
    "MACARON PETIT VERT": "39",
    "MACARON PETIT JAUNE": "39",
    "MACARON PETIT BLEU": "39",
    "MACARON GR ROUGE": "39",
    "MACARON GR CHOCOLAT NOIRE": "39",
    "MACARON PETIT ROSE": "39",
    "MACARON PETIT VIOLET": "39",
    "TROMPE L OEIL MANGUE": "40",
    "TROMPE L OEIL CITRON": "41",
    "TROMPE L OEIL CACAHUETE": "42",
    "TROMPE L OEIL FRAMBOISE": "43",
    "TROMPE L OEIL CAFE": "44",
    "TROMPE L OEIL PINACOLADA": "45",
    "TROMPE L OEIL NOISETTE": "46",
    "TABLETTE KOUNAFA": "47",
    "GATEAU CASABIANCA": "50",
    "SOIREE BLANC": "21",
    "SOIREE CITRON": "21",
    "SOIREE NOIRE": "21",
    "MILLE FEUILLE VANILLE": "22",
    "MILLE FEUILLE CHOCOLAT": "22",
    "CHAUSSON JALOUSIE AMANDE": "15",
    "PALMIER BAGUETTE SUCRE": "16",
}


def canonical_product_name(value: str) -> str:
    cleaned = clean(value)
    return ALIASES.get(key(cleaned), cleaned)


def canonical_raw_name(value: str) -> str:
    cleaned = clean(value)
    return RAW_ALIASES.get(key(cleaned), cleaned)


def product_key(product_type: str, name: str) -> str:
    return ("SF::" if product_type == "semi_finished" else "PF::") + key(name)


def product_name_key(internal_key: str) -> str:
    return internal_key.split("::", 1)[1]


def sql_text(value: str | None) -> str:
    if value is None:
        return "null"
    return "'" + value.replace("'", "''") + "'"


def infer_unit(name: str) -> str:
    normalized = key(name)
    if normalized in {"OEUF", "OEUFS", "BLANC D OEUF", "JAUNE D OEUF"}:
        return "piece"
    if normalized == "EAU" or normalized.startswith("LAIT") or normalized.startswith("HUILE") or normalized.startswith("JUS "):
        return "L"
    return "kg"


def parse_codification() -> tuple[dict[str, str], set[str]]:
    document = Document(CODIFICATION_FILE)
    code_by_name: dict[str, str] = {}
    known_semi_finished: set[str] = set()
    for table_index, table in enumerate(document.tables, start=1):
        for row in table.rows[1:]:
            values = [clean(cell.text) for cell in row.cells]
            if len(values) < 3 or not values[0] or not values[2]:
                continue
            code_by_name.setdefault(key(values[0]), values[2].replace(" …++++", "").strip())
            if table_index == 4 or "SF" in values[1].upper():
                known_semi_finished.add(key(values[0]))

    lot_code_source = (OUTPUT_DIR.parent.parent / "src" / "lib" / "productionLotCodification.ts").read_text(encoding="utf-8")
    for match in re.finditer(r'^\s*(?:"([^"]+)"|([A-Z][A-Z0-9 ]*)):\s*"([^"]+)"', lot_code_source, re.MULTILINE):
        name = match.group(1) or match.group(2)
        code_by_name.setdefault(key(name), match.group(3))
    return code_by_name, known_semi_finished


def target_type(header: str) -> str | None:
    normalized = key(header)
    if "SEMI" in normalized or normalized in {"PRODUIT SF", "PSF"}:
        return "semi_finished"
    if "FINI" in normalized or normalized == "PF":
        return "finished"
    return None


def component_columns(headers: list[str]) -> tuple[list[int], list[int]]:
    semi_columns: list[int] = []
    raw_columns: list[int] = []
    for index, header in enumerate(headers[1:], start=1):
        normalized = key(header)
        if any(token in normalized for token in ("LOT", "DF", "DATE", "OBSERVATION", "ACTION")):
            continue
        if "PRODUIT SF" in normalized or "SEMI" in normalized or normalized in {"PSF", "P SEMI FINI"}:
            semi_columns.append(index)
        elif normalized == "MP":
            raw_columns.append(index)
    return semi_columns, raw_columns


def parse_recipes(known_semi_finished: set[str]):
    products: dict[str, dict[str, str | None]] = {}
    schemas: dict[str, list[str]] = defaultdict(list)
    component_candidates: list[str] = []

    for file_name, default_category in RECIPE_FILES.items():
        document = Document(DOWNLOADS / file_name)
        for table in document.tables:
            rows = [[clean(cell.text) for cell in row.cells] for row in table.rows]
            if not rows:
                continue
            parsed_type = target_type(rows[0][0])
            if parsed_type is None:
                continue
            headers = rows[0]
            semi_columns, raw_columns = component_columns(headers)

            for row in rows[1:]:
                if not row or not row[0]:
                    continue
                target_name = canonical_product_name(row[0])
                target_key = key(target_name)
                if target_key in {"PRODUIT SF", "PRODUIT SEMI FINI", "PRODUIT FINI", "PF", "PSF"}:
                    continue
                product_type = TYPE_OVERRIDES.get(target_key, parsed_type)
                category = CATEGORY_OVERRIDES.get(target_key, default_category)
                internal_target_key = product_key(product_type, target_name)
                products.setdefault(
                    internal_target_key,
                    {
                        "name": target_name,
                        "type": product_type,
                        "category": category,
                        "official_code": None,
                    },
                )
                semi_values = [canonical_product_name(row[index]) for index in semi_columns if index < len(row) and row[index]]
                raw_values = [row[index] for index in raw_columns if index < len(row) and row[index]]

                if product_type == "finished" and semi_values:
                    for semi_name in semi_values:
                        if key(semi_name) != target_key:
                            schemas[internal_target_key].append(semi_name)
                    for raw_value in raw_values:
                        for semi_name in semi_values:
                            schemas[product_key("semi_finished", semi_name)].append(raw_value)
                else:
                    for component in semi_values + raw_values:
                        if key(canonical_product_name(component)) != target_key:
                            schemas[internal_target_key].append(component)
                component_candidates.extend(semi_values + raw_values)

    # Products explicitly listed as semi-finished in the codification are never raw materials.
    for component in component_candidates:
        canonical = canonical_product_name(component)
        normalized = key(canonical)
        internal_component_key = product_key("semi_finished", canonical)
        if internal_component_key in products:
            continue
        if normalized in known_semi_finished:
            products[internal_component_key] = {
                "name": canonical,
                "type": "semi_finished",
                "category": CATEGORY_OVERRIDES.get(normalized, "patisserie"),
                "official_code": None,
            }

    return products, schemas


def apply_context_aliases(products: dict[str, dict[str, str | None]], schemas: dict[str, list[str]]):
    contextual = {
        ("CAKE AMR FERRERO", "ENTREMET CAM"): "Entremet cake American Ferrero",
        ("CAKE AMR NOUGAT", "ENTREMET CAM"): "Entremet cake American Nougat",
        ("CAKE AMR NUTELLA", "ENTREMET CAM"): "Entremet cake American Nutella",
        ("CAKE AMR CAROTTE", "ENTREMET CAM CAROTTE"): "Entremet cake American carotte",
        ("TROMPE L OEIL CITRON", "MOUSSE"): "Mousse citron",
        ("SILICONE CITRON", "MOUSSE"): "Mousse citron",
    }
    for internal_target_key, components in list(schemas.items()):
        target_key = product_name_key(internal_target_key)
        resolved = []
        for component in components:
            replacement = contextual.get((target_key, key(component)), canonical_product_name(component))
            resolved.append(replacement)
        schemas[internal_target_key] = resolved


def resolve_products_and_components(products, schemas, code_by_name, known_semi_finished):
    semi_finished_by_name = {
        product_name_key(internal_key): internal_key
        for internal_key, product in products.items()
        if product["type"] == "semi_finished"
    }

    for internal_key, product in products.items():
        name_key = product_name_key(internal_key)
        product["official_code"] = OFFICIAL_CODE_OVERRIDES.get(name_key) or code_by_name.get(name_key)

    resolved_schemas: dict[str, list[str]] = {}
    raw_materials: dict[str, str] = {}
    for internal_target_key, components in schemas.items():
        if internal_target_key not in products:
            continue
        unique: dict[str, str] = {}
        for component_value in components:
            manufactured_name = canonical_product_name(component_value)
            component_key = key(manufactured_name)
            if component_key in semi_finished_by_name or component_key in known_semi_finished:
                internal_component_key = semi_finished_by_name.get(component_key)
                component_name = products.get(internal_component_key, {}).get("name") if internal_component_key else manufactured_name
            else:
                component_name = canonical_raw_name(component_value)
                component_key = key(component_name)
                raw_materials.setdefault(component_key, component_name)
            if component_key != product_name_key(internal_target_key):
                unique.setdefault(component_key, str(component_name))
        if unique:
            resolved_schemas[internal_target_key] = list(unique.values())

    # Any manufactured component recognized from codification but absent from the parsed targets.
    for components in resolved_schemas.values():
        for component_name in components:
            component_key = key(component_name)
            internal_component_key = product_key("semi_finished", component_name)
            if component_key in raw_materials or internal_component_key in products:
                continue
            products[internal_component_key] = {
                "name": component_name,
                "type": "semi_finished",
                "category": CATEGORY_OVERRIDES.get(component_key, "patisserie"),
                "official_code": OFFICIAL_CODE_OVERRIDES.get(component_key) or code_by_name.get(component_key),
            }

    # Every schema must use the exact canonical spelling selected for its raw material.
    for target_key, components in resolved_schemas.items():
        resolved_schemas[target_key] = [raw_materials.get(key(component), component) for component in components]
    return resolved_schemas, raw_materials


def detect_cycles(products, schemas):
    graph = {
        target: [product_key("semi_finished", component) for component in components if product_key("semi_finished", component) in products]
        for target, components in schemas.items()
    }
    visiting: set[str] = set()
    visited: set[str] = set()

    def visit(node: str, path: list[str]):
        if node in visiting:
            raise RuntimeError("Circular schema detected: " + " -> ".join(path + [node]))
        if node in visited:
            return
        visiting.add(node)
        for child in graph.get(node, []):
            visit(child, path + [node])
        visiting.remove(node)
        visited.add(node)

    for internal_key in graph:
        visit(internal_key, [])


def values_sql(rows: list[tuple[str, ...]]) -> str:
    return ",\n".join("  (" + ", ".join(sql_text(value) for value in row) + ")" for row in rows)


def raw_material_sql(raw_materials: dict[str, str]) -> str:
    rows = [(name, infer_unit(name)) for name in sorted(raw_materials.values(), key=str.casefold)]
    return f"""-- Self-contained and safe to run more than once.
with material_input(name, unit) as (
  values
{values_sql(rows)}
),
inserted_supplier as (
  insert into suppliers (name, is_active)
  select 'DIVERS', true
  where not exists (select 1 from suppliers where lower(trim(name)) = lower('DIVERS'))
  returning id
),
divers_supplier as (
  select id from inserted_supplier
  union all
  select id from suppliers
  where lower(trim(name)) = lower('DIVERS')
    and not exists (select 1 from inserted_supplier)
  order by id
  limit 1
),
inserted_products as (
  insert into products (code, name, type, unit)
  select
    'MP-DIVERS-' || upper(substr(md5(input.name), 1, 12)),
    input.name,
    'raw'::product_type,
    input.unit
  from material_input input
  where not exists (
    select 1 from products
    where type = 'raw' and lower(trim(products.name)) = lower(trim(input.name))
  )
  on conflict (code) do nothing
  returning id, name
),
all_materials as (
  select inserted_products.id, inserted_products.name
  from inserted_products
  union all
  select products.id, products.name
  from products
  join material_input input on lower(trim(input.name)) = lower(trim(products.name))
  where products.type = 'raw' and products.is_active = true
    and not exists (
      select 1 from inserted_products
      where lower(trim(inserted_products.name)) = lower(trim(products.name))
    )
),
inserted_links as (
  insert into supplier_raw_materials (supplier_id, product_id)
  select divers_supplier.id, all_materials.id
  from divers_supplier cross join all_materials
  on conflict (supplier_id, product_id) do nothing
  returning product_id
)
select
  (select count(*) from material_input) as requested_raw_materials,
  (select count(*) from inserted_products) as newly_created_raw_materials,
  (select count(*) from inserted_links) as newly_linked_to_divers,
  (select count(*) from all_materials) as total_import_materials,
  (select jsonb_agg(jsonb_build_object('name', name, 'unit', unit) order by name) from material_input) as raw_materials;
"""


def input_payloads(products, schemas) -> tuple[str, str]:
    product_records = []
    for product in sorted(products.values(), key=lambda item: (str(item["name"]).casefold(), str(item["type"]))):
        product_records.append(
            {
                "name": str(product["name"]),
                "product_type": str(product["type"]),
                "category": str(product["category"]),
                "official_code": product["official_code"],
                "unit": "unites" if product["type"] == "finished" else "kg",
            }
        )
    link_records = []
    for internal_target_key, components in sorted(schemas.items(), key=lambda item: str(products[item[0]]["name"]).casefold()):
        target = products[internal_target_key]
        for order, component in enumerate(components, start=1):
            link_records.append(
                {
                    "target_name": str(target["name"]),
                    "target_type": str(target["type"]),
                    "target_category": str(target["category"]),
                    "component_name": component,
                    "component_order": order,
                }
            )
    return (
        json.dumps(product_records, ensure_ascii=False, separators=(",", ":")),
        json.dumps(link_records, ensure_ascii=False, separators=(",", ":")),
    )


def product_schema_sql(products, schemas) -> str:
    product_json, link_json = input_payloads(products, schemas)
    return f"""-- Self-contained. Run after 01_insert_divers_raw_materials.sql.
do $import$
declare
  v_products jsonb := $products${product_json}$products$::jsonb;
  v_links jsonb := $links${link_json}$links$::jsonb;
  missing_components text[];
  ambiguous_names text[];
  schema_record record;
  target_id uuid;
  component_ids uuid[];
  v_diagram_nodes jsonb;
  v_diagram_edges jsonb;
  v_recipe_id uuid;
  v_schema_count integer;
  v_missing_code_count integer;
begin
  select array_agg(distinct input.component_name order by input.component_name)
  into missing_components
  from jsonb_to_recordset(v_links) as input(
    target_name text, target_type text, target_category text, component_name text, component_order integer
  )
  where not exists (
    select 1 from products
    where is_active = true
      and type in ('raw', 'semi_finished')
      and lower(trim(products.name)) = lower(trim(input.component_name))
  )
  and not exists (
    select 1
    from jsonb_to_recordset(v_products) as product_input(
      name text, product_type text, category text, official_code text, unit text
    )
    where product_input.product_type = 'semi_finished'
      and lower(trim(product_input.name)) = lower(trim(input.component_name))
  );

  if coalesce(array_length(missing_components, 1), 0) > 0 then
    raise exception 'Run 01_insert_divers_raw_materials.sql first. Missing components: %', missing_components;
  end if;

  update products
  set type = input.product_type::product_type,
      category = input.category,
      unit = input.unit,
      code = (case when input.product_type = 'finished' then 'PF-' else 'SF-' end)
        || upper(input.category) || '-'
        || case
          when nullif(regexp_replace(input.official_code, '[^A-Za-z0-9-]+', '', 'g'), '') is not null
            then 'CODE-' || regexp_replace(input.official_code, '[^A-Za-z0-9-]+', '', 'g')
          else 'AUTO-' || upper(substr(md5(input.name), 1, 8))
        end
        || '-' || upper(substr(md5(input.name), 1, 6)),
      updated_at = now()
  from jsonb_to_recordset(v_products) as input(
    name text, product_type text, category text, official_code text, unit text
  )
  where products.type = 'raw'
    and lower(trim(products.name)) = lower(trim(input.name))
    and not exists (
      select 1 from products manufactured
      where manufactured.type = input.product_type::product_type
        and manufactured.category = input.category
        and lower(trim(manufactured.name)) = lower(trim(input.name))
    );

  update products
  set unit = input.unit, is_active = true, updated_at = now()
  from jsonb_to_recordset(v_products) as input(
    name text, product_type text, category text, official_code text, unit text
  )
  where products.type = input.product_type::product_type
    and products.category = input.category
    and lower(trim(products.name)) = lower(trim(input.name));

  insert into products (code, name, type, category, unit)
  select
    (case when input.product_type = 'finished' then 'PF-' else 'SF-' end)
      || upper(input.category) || '-'
      || case
        when nullif(regexp_replace(input.official_code, '[^A-Za-z0-9-]+', '', 'g'), '') is not null
          then 'CODE-' || regexp_replace(input.official_code, '[^A-Za-z0-9-]+', '', 'g')
        else 'AUTO-' || upper(substr(md5(input.name), 1, 8))
      end
      || '-' || upper(substr(md5(input.name), 1, 6)),
    input.name,
    input.product_type::product_type,
    input.category,
    input.unit
  from jsonb_to_recordset(v_products) as input(
    name text, product_type text, category text, official_code text, unit text
  )
  where not exists (
    select 1 from products
    where products.type = input.product_type::product_type
      and products.category = input.category
      and lower(trim(products.name)) = lower(trim(input.name))
  )
  on conflict (code) do nothing;

  select array_agg(name || ' (' || matches || ' matches)' order by name)
  into ambiguous_names
  from (
    select input.name, input.product_type, input.category, count(products.id) as matches
    from jsonb_to_recordset(v_products) as input(
      name text, product_type text, category text, official_code text, unit text
    )
    left join products on products.type = input.product_type::product_type
      and products.category = input.category
      and lower(trim(products.name)) = lower(trim(input.name))
    group by input.name, input.product_type, input.category
    having count(products.id) <> 1
  ) problems;

  if coalesce(array_length(ambiguous_names, 1), 0) > 0 then
    raise exception 'Missing or ambiguous imported products: %', ambiguous_names;
  end if;

  select array_agg(target_name || ' -> ' || component_name order by target_name, component_name)
  into ambiguous_names
  from (
    select distinct input.target_name, input.target_type, input.target_category, input.component_name
    from jsonb_to_recordset(v_links) as input(
      target_name text, target_type text, target_category text, component_name text, component_order integer
    )
    where not exists (
      select 1 from products
      where products.is_active = true
        and lower(trim(products.name)) = lower(trim(input.component_name))
        and (
          products.type = 'raw'
          or (products.type = 'semi_finished' and products.category = input.target_category)
          or (
            products.type = 'semi_finished'
            and 1 = (
              select count(*) from products candidate
              where candidate.type = 'semi_finished'
                and candidate.is_active = true
                and lower(trim(candidate.name)) = lower(trim(input.component_name))
            )
          )
        )
    )
  ) problems;

  if coalesce(array_length(ambiguous_names, 1), 0) > 0 then
    raise exception 'Missing or ambiguous schema components: %', ambiguous_names;
  end if;

  for schema_record in
    select input.target_name, input.target_type, input.target_category
    from jsonb_to_recordset(v_links) as input(
      target_name text, target_type text, target_category text, component_name text, component_order integer
    )
    group by input.target_name, input.target_type, input.target_category
    order by case when input.target_type = 'semi_finished' then 0 else 1 end, input.target_name
  loop
    select id into target_id
    from products
    where type = schema_record.target_type::product_type
      and category = schema_record.target_category
      and lower(trim(name)) = lower(trim(schema_record.target_name));

    select array_agg(product_id order by component_order)
    into component_ids
    from (
      select distinct on (lower(trim(input.component_name)))
        input.component_order,
        resolved_product.id as product_id
      from jsonb_to_recordset(v_links) as input(
        target_name text, target_type text, target_category text, component_name text, component_order integer
      )
      join lateral (
        select candidate.id
        from products candidate
        where candidate.type in ('raw', 'semi_finished')
          and candidate.is_active = true
          and lower(trim(candidate.name)) = lower(trim(input.component_name))
        order by
          case
            when candidate.type = 'semi_finished' and candidate.category = schema_record.target_category then 0
            when candidate.type = 'raw' then 1
            else 2
          end,
          candidate.created_at
        limit 1
      ) resolved_product on true
      where lower(trim(input.target_name)) = lower(trim(schema_record.target_name))
        and input.target_type = schema_record.target_type
        and input.target_category = schema_record.target_category
      order by lower(trim(input.component_name)), input.component_order
    ) resolved;

    with recursive diagram_tree as (
      select
        component_id as product_id,
        target_id::text as parent_node_id,
        target_id::text || '__' || component_id::text as node_id,
        1 as depth,
        ordinality::integer as component_order,
        lpad(ordinality::text, 4, '0') as path_sort,
        array[target_id, component_id] as path
      from unnest(component_ids) with ordinality as components(component_id, ordinality)
      union all
      select
        resolved_product.id as product_id,
        diagram_tree.node_id as parent_node_id,
        diagram_tree.node_id || '__' || resolved_product.id::text as node_id,
        diagram_tree.depth + 1 as depth,
        child_input.component_order,
        diagram_tree.path_sort || '.' || lpad(child_input.component_order::text, 4, '0') as path_sort,
        diagram_tree.path || resolved_product.id
      from diagram_tree
      join products parent_product on parent_product.id = diagram_tree.product_id
        and parent_product.type = 'semi_finished'
        and parent_product.is_active = true
      join jsonb_to_recordset(v_links) as child_input(
        target_name text, target_type text, target_category text, component_name text, component_order integer
      ) on lower(trim(child_input.target_name)) = lower(trim(parent_product.name))
        and child_input.target_type = parent_product.type::text
        and child_input.target_category = parent_product.category
      join lateral (
        select candidate.id
        from products candidate
        where candidate.type in ('raw', 'semi_finished')
          and candidate.is_active = true
          and lower(trim(candidate.name)) = lower(trim(child_input.component_name))
        order by
          case
            when candidate.type = 'semi_finished' and candidate.category = parent_product.category then 0
            when candidate.type = 'raw' then 1
            else 2
          end,
          candidate.created_at
        limit 1
      ) resolved_product on not resolved_product.id = any(diagram_tree.path)
    ),
    ordered_nodes as (
      select
        product_id,
        parent_node_id,
        node_id,
        depth,
        row_number() over (order by path_sort, product_id::text) as row_index,
        count(*) over () as total_nodes
      from diagram_tree
    )
    select
      jsonb_build_array(
        jsonb_build_object(
          'id', target_id::text,
          'type', 'product',
          'position', jsonb_build_object('x', -320, 'y', 0),
          'data', jsonb_build_object('productId', target_id::text, 'isTarget', true)
        )
      ) || coalesce(
        jsonb_agg(
          jsonb_build_object(
            'id', node_id,
            'type', 'product',
            'position', jsonb_build_object(
              'x', ((depth - 1) * 320),
              'y', ((row_index - ((total_nodes::numeric + 1) / 2)) * 165)::integer
            ),
            'data', jsonb_build_object('productId', product_id::text, 'isTarget', false)
          )
          order by row_index
        ),
        '[]'::jsonb
      ),
      coalesce(
        jsonb_agg(
          jsonb_build_object(
            'id', parent_node_id || '->' || node_id,
            'source', parent_node_id,
            'target', node_id,
            'type', 'smoothstep'
          )
          order by row_index
        ),
        '[]'::jsonb
      )
    into v_diagram_nodes, v_diagram_edges
    from ordered_nodes;

    select id into v_recipe_id
    from recipes
    where product_id = target_id and is_active = true
    order by version desc
    limit 1;

    if v_recipe_id is null then
      insert into recipes (product_id, version, is_active, notes, diagram_nodes, diagram_edges, diagram_viewport)
      values (
        target_id,
        coalesce((select max(version) from recipes where product_id = target_id), 0) + 1,
        true,
        'Casabianca recipe document import',
        v_diagram_nodes,
        v_diagram_edges,
        jsonb_build_object('x', 0, 'y', 0, 'zoom', 0.9)
      )
      returning id into v_recipe_id;
    else
      update recipes
      set notes = 'Casabianca recipe document import',
          diagram_nodes = v_diagram_nodes,
          diagram_edges = v_diagram_edges,
          diagram_viewport = jsonb_build_object('x', 0, 'y', 0, 'zoom', 0.9)
      where id = v_recipe_id;
    end if;

    update recipes set is_active = false
    where product_id = target_id and id <> v_recipe_id and is_active = true;

    delete from recipe_components where recipe_components.recipe_id = v_recipe_id;

    insert into recipe_components (recipe_id, component_product_id, quantity, unit)
    select v_recipe_id, component_id, null, null
    from unnest(component_ids) component_id;

    update products set updated_at = now() where id = target_id;
  end loop;

  select count(*) into v_schema_count
  from (
    select distinct input.target_name, input.target_type, input.target_category
    from jsonb_to_recordset(v_links) as input(
      target_name text, target_type text, target_category text, component_name text, component_order integer
    )
  ) schemas;

  select count(*) into v_missing_code_count
  from jsonb_to_recordset(v_products) as input(
    name text, product_type text, category text, official_code text, unit text
  )
  where input.official_code is null;

  raise notice 'Imported % products, % schemas, % component links. % products have no official codification.',
    jsonb_array_length(v_products), v_schema_count, jsonb_array_length(v_links), v_missing_code_count;
end
$import$;
"""


def verification_sql(products, schemas) -> str:
    product_json, link_json = input_payloads(products, schemas)
    return f"""-- Self-contained verification query.
with expected_products as (
  select input.name, input.product_type::product_type as product_type, input.category, input.official_code, input.unit
  from jsonb_to_recordset($products${product_json}$products$::jsonb) as input(
    name text, product_type text, category text, official_code text, unit text
  )
),
schema_link_input as (
  select input.target_name, input.target_type::product_type as target_type,
    input.target_category, input.component_name, input.component_order
  from jsonb_to_recordset($links${link_json}$links$::jsonb) as input(
    target_name text, target_type text, target_category text, component_name text, component_order integer
  )
),
actual_products as (
  select input.name, input.product_type, products.id, products.category, products.unit
  from expected_products input
  join products on products.type = input.product_type
    and products.category = input.category
    and products.is_active = true
    and lower(trim(products.name)) = lower(trim(input.name))
),
missing_products as (
  select input.name
  from expected_products input
  where not exists (
    select 1 from actual_products
    where lower(trim(actual_products.name)) = lower(trim(input.name))
      and actual_products.product_type = input.product_type
      and actual_products.category = input.category
  )
),
actual_links as (
  select target.name as target_name, target.type as target_type, target.category as target_category,
    component.name as component_name, recipes.id as recipe_id
  from recipes
  join products target on target.id = recipes.product_id
  join recipe_components link on link.recipe_id = recipes.id
  join products component on component.id = link.component_product_id
  where recipes.is_active = true
    and exists (
      select 1 from schema_link_input input
      where lower(trim(input.target_name)) = lower(trim(target.name))
        and input.target_type = target.type
        and input.target_category = target.category
    )
),
missing_links as (
  select input.target_name, input.component_name
  from schema_link_input input
  where not exists (
    select 1 from actual_links
    where lower(trim(actual_links.target_name)) = lower(trim(input.target_name))
      and actual_links.target_type = input.target_type
      and actual_links.target_category = input.target_category
      and lower(trim(actual_links.component_name)) = lower(trim(input.component_name))
  )
),
extra_links as (
  select actual.target_name, actual.component_name
  from actual_links actual
  where not exists (
    select 1 from schema_link_input input
    where lower(trim(input.target_name)) = lower(trim(actual.target_name))
      and input.target_type = actual.target_type
      and input.target_category = actual.target_category
      and lower(trim(input.component_name)) = lower(trim(actual.component_name))
  )
),
duplicate_active_recipes as (
  select target_name, target_type, target_category, count(distinct recipe_id) as active_count
  from actual_links
  group by target_name, target_type, target_category
  having count(distinct recipe_id) > 1
),
recipes_with_diagrams as (
  select distinct target.name as target_name, target.type as target_type, target.category as target_category
  from recipes
  join products target on target.id = recipes.product_id
  where recipes.is_active = true
    and jsonb_array_length(coalesce(recipes.diagram_nodes, '[]'::jsonb)) > 0
    and jsonb_array_length(coalesce(recipes.diagram_edges, '[]'::jsonb)) > 0
    and exists (
      select 1 from schema_link_input input
      where lower(trim(input.target_name)) = lower(trim(target.name))
        and input.target_type = target.type
        and input.target_category = target.category
    )
),
missing_diagrams as (
  select distinct input.target_name
  from schema_link_input input
  where not exists (
    select 1 from recipes_with_diagrams diagrams
    where lower(trim(diagrams.target_name)) = lower(trim(input.target_name))
      and diagrams.target_type = input.target_type
      and diagrams.target_category = input.target_category
  )
),
missing_schema_targets as (
  select distinct input.target_name
  from schema_link_input input
  where not exists (
    select 1 from actual_links where lower(trim(actual_links.target_name)) = lower(trim(input.target_name))
      and actual_links.target_type = input.target_type
      and actual_links.target_category = input.target_category
  )
),
unlinked_divers_materials as (
  select component.name
  from products component
  where component.type = 'raw'
    and component.is_active = true
    and exists (
      select 1 from schema_link_input input
      where lower(trim(input.component_name)) = lower(trim(component.name))
    )
    and not exists (
      select 1
      from supplier_raw_materials link
      join suppliers supplier on supplier.id = link.supplier_id
      where link.product_id = component.id and lower(trim(supplier.name)) = lower('DIVERS')
    )
)
select
  (select count(*) from expected_products) as expected_products,
  (select count(*) from actual_products) as found_products,
  (select count(distinct (target_name, target_type, target_category)) from schema_link_input) as expected_schemas,
  (select count(distinct (target_name, target_type, target_category)) from actual_links) as active_schemas,
  (select count(distinct (target_name, target_type, target_category)) from recipes_with_diagrams) as active_schema_diagrams,
  (select count(*) from schema_link_input) as expected_component_links,
  (select count(*) from actual_links) as active_component_links,
  coalesce((select array_agg(name order by name) from missing_products), array[]::text[]) as missing_products,
  coalesce((select array_agg(target_name order by target_name) from missing_schema_targets), array[]::text[]) as missing_schemas,
  coalesce((select array_agg(target_name order by target_name) from missing_diagrams), array[]::text[]) as missing_diagrams,
  coalesce((select array_agg(target_name || ' -> ' || component_name order by target_name, component_name) from missing_links), array[]::text[]) as missing_component_links,
  coalesce((select array_agg(target_name || ' -> ' || component_name order by target_name, component_name) from extra_links), array[]::text[]) as extra_component_links,
  coalesce((select array_agg(target_name || ' (' || active_count || ' active)' order by target_name) from duplicate_active_recipes), array[]::text[]) as duplicate_active_recipes,
  coalesce((select array_agg(name order by name) from unlinked_divers_materials), array[]::text[]) as raw_materials_not_linked_to_divers,
  coalesce((select array_agg(name order by name) from expected_products where official_code is null), array[]::text[]) as products_without_official_codification;
"""


def divers_link_repair_sql(products, schemas) -> str:
    _, link_json = input_payloads(products, schemas)
    return f"""-- Repairs DIVERS supplier links for raw materials used by the Casabianca import.
-- Safe to run more than once.
with schema_link_input as (
  select input.component_name
  from jsonb_to_recordset($links${link_json}$links$::jsonb) as input(
    target_name text, target_type text, target_category text, component_name text, component_order integer
  )
),
divers_supplier as (
  select id
  from suppliers
  where lower(trim(name)) = lower('DIVERS')
  order by id
  limit 1
),
raw_components as (
  select distinct raw_product.id, raw_product.name
  from schema_link_input input
  join products raw_product on raw_product.type = 'raw'
    and raw_product.is_active = true
    and lower(trim(raw_product.name)) = lower(trim(input.component_name))
),
inserted_links as (
  insert into supplier_raw_materials (supplier_id, product_id)
  select divers_supplier.id, raw_components.id
  from divers_supplier
  cross join raw_components
  where divers_supplier.id is not null
  on conflict (supplier_id, product_id) do nothing
  returning product_id
),
remaining_unlinked as (
  select raw_components.name
  from raw_components
  where not exists (
    select 1
    from supplier_raw_materials link
    join divers_supplier on divers_supplier.id = link.supplier_id
    where link.product_id = raw_components.id
  )
  and not exists (
    select 1
    from inserted_links
    where inserted_links.product_id = raw_components.id
  )
)
select
  case when exists (select 1 from divers_supplier) then 'DIVERS supplier found' else 'DIVERS supplier missing' end as supplier_status,
  (select count(*) from raw_components) as raw_components_used_by_import,
  (select count(*) from inserted_links) as newly_linked_to_divers,
  coalesce((select array_agg(name order by name) from remaining_unlinked), array[]::text[]) as remaining_unlinked_raw_materials;
"""


def main():
    missing_files = [name for name in RECIPE_FILES if not (DOWNLOADS / name).exists()]
    if missing_files:
        raise FileNotFoundError("Missing recipe documents: " + ", ".join(missing_files))
    if not CODIFICATION_FILE.exists():
        raise FileNotFoundError(f"Missing codification document: {CODIFICATION_FILE}")

    code_by_name, known_semi_finished = parse_codification()
    known_semi_finished.update(key(name) for name in FORCED_SEMI_FINISHED)
    products, schemas = parse_recipes(known_semi_finished)
    apply_context_aliases(products, schemas)
    resolved_schemas, raw_materials = resolve_products_and_components(products, schemas, code_by_name, known_semi_finished)
    detect_cycles(products, resolved_schemas)

    outputs = {
        "01_insert_divers_raw_materials.sql": raw_material_sql(raw_materials),
        "02_insert_casabianca_products_and_schemas.sql": product_schema_sql(products, resolved_schemas),
        "03_verify_casabianca_products_and_schemas.sql": verification_sql(products, resolved_schemas),
        "04_link_casabianca_raw_materials_to_divers.sql": divers_link_repair_sql(products, resolved_schemas),
    }
    for file_name, contents in outputs.items():
        (OUTPUT_DIR / file_name).write_text(contents, encoding="utf-8", newline="\n")

    missing_codes = sorted(
        str(product["name"])
        for product in products.values()
        if product["official_code"] is None
    )
    print(f"Generated {len(raw_materials)} raw materials, {len(products)} products, {len(resolved_schemas)} schemas, and {sum(map(len, resolved_schemas.values()))} component links.")
    print(f"Products without an official codification: {len(missing_codes)}")
    for name in missing_codes:
        print(f"  - {name}")


if __name__ == "__main__":
    main()
